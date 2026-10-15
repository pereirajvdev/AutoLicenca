import argparse
import time
from pathlib import Path

import pyautogui
import pyperclip
import win32gui

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime


# ============================================================
# CONFIGURAÇÕES
# ============================================================

TITULO_MENU = "Menu: Arquivos / Cadastro de Funcionários"

X_INICIAL = 583
Y_INICIAL = 841

X_LICENCAS = 667
Y_LICENCAS = 495

X_PROXIMA_ACAO = 898
Y_PROXIMA_ACAO = 419

INTERVALO = 0.3
INTERVALO_VERIFICACAO = 0.5


# ============================================================
# IDENTIFICAÇÃO DE JANELAS
# ============================================================

def encontrar_janela_por_titulo(titulo_procurado):
    """
    Retorna o identificador da primeira janela visível
    cujo título contenha o texto procurado.
    """
    hwnd_encontrado = None
    titulo_procurado = titulo_procurado.casefold()

    def callback(hwnd, _):
        nonlocal hwnd_encontrado

        try:
            titulo = win32gui.GetWindowText(hwnd).strip()

            if (
                win32gui.IsWindowVisible(hwnd)
                and titulo_procurado in titulo.casefold()
            ):
                hwnd_encontrado = hwnd
                return False

        except Exception:
            pass

        return True

    win32gui.EnumWindows(callback, None)

    return hwnd_encontrado


def aguardar_estado_menu(timeout=30):
    """
    Aguarda até que o Menu esteja aberto.
    """
    limite = time.time() + timeout

    while time.time() < limite:
        if encontrar_janela_por_titulo(TITULO_MENU) is not None:
            return True

        time.sleep(INTERVALO_VERIFICACAO)

    return False


def ativar_janela(titulo, timeout=15):
    """
    Localiza e ativa a janela pelo título.
    """
    print(f"Esperando janela: {titulo}")

    limite = time.time() + timeout

    while time.time() < limite:
        hwnd = encontrar_janela_por_titulo(titulo)

        if hwnd:
            try:
                win32gui.SetForegroundWindow(hwnd)
                time.sleep(INTERVALO)

                print(f"Janela ativa: {win32gui.GetWindowText(hwnd)}")
                return True

            except Exception as erro:
                print(f"Não foi possível ativar a janela: {erro}")
                return False

        time.sleep(0.2)

    print(f"Janela não encontrada: {titulo}")
    return False


# ============================================================
# LEITURA DA PLANILHA
# ============================================================

def ler_servidores(caminho_planilha):
    """
    Lê exclusivamente os nomes dos servidores na coluna B.
    Ignora completamente a coluna A (matrículas).
    """
    workbook = load_workbook(
        caminho_planilha,
        read_only=True,
        data_only=True,
    )

    try:
        planilha = workbook.active
        servidores = []

        for linha in planilha.iter_rows(
            min_col=2,
            max_col=2,
            values_only=True,
        ):
            valor = linha[0]

            if valor is None:
                continue

            nome = str(valor).strip()

            if nome:
                servidores.append(nome)

        return servidores

    finally:
        workbook.close()


# ============================================================
# PESQUISA DO SERVIDOR
# ============================================================

def pesquisar_servidor(nome, numero_linha=None, caminho_planilha=None):
    print(f"\nPesquisando servidor: {nome}")
    """
    Executa a sequência inicial:
    1. Aguarda o Menu.
    2. Ativa a janela.
    3. Clica na posição inicial.
    4. Cola o nome.
    5. Pressiona Enter duas vezes.
    """

    print(f"\n{'=' * 50}")
    print(f"Linha {numero_linha}: {nome}")

    print("Aguardando o Menu...")

    if not aguardar_estado_menu():
        print("Menu não encontrado. Servidor ignorado.")
        return False

    print(f"Janela encontrada: {TITULO_MENU}")

    if not ativar_janela(TITULO_MENU):
        return False

    print("Clicando na posição inicial...")
    pyautogui.click(X_INICIAL, Y_INICIAL)
    time.sleep(INTERVALO)

    print(f"Pesquisando servidor: {nome}")
    pyperclip.copy(nome)
    pyautogui.hotkey("ctrl", "v")
    time.sleep(INTERVALO)

    print("Pressionando primeiro Enter...")
    pyautogui.press("enter")
    time.sleep(INTERVALO)

    print("Pressionando segundo Enter...")
    pyautogui.press("enter")

    print("Abrindo janela de Licenças...")
    pyautogui.click(X_LICENCAS, Y_LICENCAS)
    time.sleep(INTERVALO)

    print("Executando próximo clique...")
    pyautogui.click(X_PROXIMA_ACAO, Y_PROXIMA_ACAO)
    time.sleep(INTERVALO)

    print("Clique para abrir Licenças realizado.")

    obter_licencas(nome, caminho_planilha)
    
    time.sleep(1)

    print(f"Sequência concluída para: {nome}")

    return True


def obter_licencas(nome_servidor, caminho_planilha):
    """
    Lê os registros da grade até encontrar duas cópias
    consecutivas iguais e salva o resultado em Excel.
    """
    titulo = "Licenças / Afastamentos / Exames"
    limite_tentativas = 100

    hwnd = encontrar_janela_por_titulo(titulo)

    if not hwnd:
        print("Janela de Licenças não encontrada.")
        return

    print(f"\nObtendo licenças de: {nome_servidor}")

    win32gui.SetForegroundWindow(hwnd)
    time.sleep(0.5)

    pyautogui.click(800, 460)
    time.sleep(0.3)

    registros = []
    anterior = None

    for tentativa in range(1, limite_tentativas + 1):
        pyperclip.copy("")
        time.sleep(0.1)

        pyautogui.hotkey("ctrl", "c")
        time.sleep(0.5)

        atual = pyperclip.paste()

        if not atual.strip():
            print("Nenhum texto foi copiado.")
            break

        # Encerra quando o conteúdo se repetir.
        if anterior is not None and atual == anterior:
            print("Fim da leitura: conteúdo repetido.")
            break

        registros.append(atual.strip())
        anterior = atual

        print(f"Registro {len(registros)} capturado.")

        pyautogui.press("down")
        time.sleep(0.3)

    if not registros:
        print("Nenhuma licença foi encontrada.")
        return

    # Salva o Excel na mesma pasta da planilha de entrada.
    caminho_planilha = Path(caminho_planilha).resolve()

    nome_arquivo = (
        f"Licencas_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    )

    caminho_saida = caminho_planilha.parent / nome_arquivo

    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Licenças"

    # Cabeçalho.
    planilha.append(["Servidor", "Nº do Registro", "Dados da Licença"])

    for celula in planilha[1]:
        celula.fill = PatternFill(
            fill_type="solid",
            fgColor="4472C4",
        )
        celula.font = Font(
            bold=True,
            color="FFFFFF",
        )
        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # Mantém cada registro integral em sua própria linha.
    for indice, registro in enumerate(registros, start=1):
        planilha.append([
            nome_servidor,
            indice,
            registro,
        ])

    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = planilha.dimensions

    larguras = {
        "A": 35,
        "B": 16,
        "C": 100,
    }

    for coluna, largura in larguras.items():
        planilha.column_dimensions[coluna].width = largura

    for linha in planilha.iter_rows(min_row=2):
        linha[0].alignment = Alignment(vertical="top")
        linha[1].alignment = Alignment(
            horizontal="center",
            vertical="top",
        )
        linha[2].alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

    workbook.save(caminho_saida)
    workbook.close()

    print("\nExcel criado com sucesso.")
    print(f"Servidor: {nome_servidor}")
    print(f"Registros: {len(registros)}")
    print(f"Arquivo: {caminho_saida.resolve()}")

        
# ============================================================
# PROGRAMA PRINCIPAL
# ============================================================

def main():
    parser = argparse.ArgumentParser(
        description="Executa a pesquisa inicial dos servidores de uma planilha."
    )

    parser.add_argument(
        "planilha",
        type=Path,
        help="Caminho da planilha Excel.",
    )

    args = parser.parse_args()

    if not args.planilha.exists():
        print(f"Planilha não encontrada: {args.planilha}")
        return

    try:
        servidores = ler_servidores(args.planilha)
    except Exception as erro:
        print(f"Erro ao ler a planilha: {erro}")
        return

    if not servidores:
        print("Nenhum nome encontrado na coluna B.")
        return

    print(f"Total de servidores encontrados: {len(servidores)}")

    for nome in servidores:
        try:
            pesquisar_servidor(
                nome,
                0,
                args.planilha,
            )
            time.sleep(0.5)

        except KeyboardInterrupt:
            print("\nProcessamento interrompido pelo usuário.")
            break

        except Exception as erro:
            print(f"Erro ao processar {nome}: {erro}")

    print("\nProcessamento finalizado.")


if __name__ == "__main__":
    main()